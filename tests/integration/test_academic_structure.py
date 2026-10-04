"""
Migration 009 and what is built on it, against a real server.

Two halves.

**The migration, run over the data it was written for.** Migrating an empty
database proves nothing (lessons.md L11), and the fixtures everywhere else in
this directory build a schema that is already at 009. So the first half builds
a *second* scratch database at 008, fills it with the shapes the deployed
database held - and the ones it did not, which are the ones to be afraid of -
and only then applies 009.

**The scope (DM-5).** An instructor sees their own classes and sessions and
nothing else; an administrator sees everything or one instructor's. Every
assertion is made through the route, because the scope is decided from the
session and a repository test cannot see that the URL is not allowed to choose
it.
"""

from __future__ import annotations

import io
import shutil

import mysql.connector
import pytest
from openpyxl import load_workbook

from config.settings import BASE_DIR, settings

from .conftest import (
    _connect_serverwide,
    make_attendance,
    make_enrolment,
    make_instructor,
    make_student,
    make_subject,
    sign_in_as,
)

pytestmark = pytest.mark.integration

MIGRATIONS = BASE_DIR / "migrations"


# ===========================================================================
# The migration, over real shapes of data
# ===========================================================================


class DatabaseAt008:
    """A scratch database at 008, a cursor into it, and the way to 009."""

    def __init__(self, name, connection):
        self.name = name
        self.connection = connection
        self.cursor = connection.cursor()

    def migrate(self):
        """Commit what the test seeded, then apply 009 from the real file."""
        from infra.migrations import apply_all

        self.connection.commit()

        apply_all(database=self.name)

        # A new snapshot: REPEATABLE READ would otherwise keep showing this
        # connection the rows as they were before the migration.
        self.connection.commit()


@pytest.fixture
def database_at_008(scratch_database, tmp_path):
    """
    A second scratch database, migrated to 008 and no further.

    Named after the first, so it inherits the `test_*` allowlist the safety
    gate in conftest.py already enforced - it cannot be the deployment.

    ⚠️ **The fixture owns the connection, and closes it before the DROP.** The
    first draft opened it in each test and closed it on the last line - so a
    failing assertion skipped the close, the open transaction kept its
    metadata lock, and `DROP DATABASE` below waited on it for ever. A test
    that *hangs* when it should fail was found by mutating the migration, which
    is the only time these tests are expected to fail.
    """
    from infra.migrations import apply_all

    name = f"{scratch_database}_at_008"

    before = tmp_path / "before"
    before.mkdir()

    for path in sorted(MIGRATIONS.glob("*.sql")):
        if path.name[:3] <= "008":
            shutil.copy(path, before / path.name)

    connection = _connect_serverwide()
    try:
        cursor = connection.cursor()
        cursor.execute(f"DROP DATABASE IF EXISTS `{name}`")
        cursor.execute(f"CREATE DATABASE `{name}` CHARACTER SET utf8mb4")
        cursor.close()
    finally:
        connection.close()

    apply_all(database=name, migrations_dir=before)

    kwargs = settings.db_kwargs()
    kwargs["database"] = name

    database = DatabaseAt008(name, mysql.connector.connect(**kwargs))

    try:
        yield database

    finally:
        database.connection.close()

        connection = _connect_serverwide()
        try:
            cursor = connection.cursor()
            cursor.execute(f"DROP DATABASE IF EXISTS `{name}`")
            cursor.close()
        finally:
            connection.close()


def legacy_student(cursor, student_id, department, program, year_level, section):
    """A student as the pre-009 forms wrote one: four unrelated strings."""
    cursor.execute(
        """
        INSERT INTO students
            (student_id, name, college_department, program, year_level, section)
        VALUES (%s, 'Integration Test Subject', %s, %s, %s, %s)
        """,
        (student_id, department, program, year_level, section),
    )


def placement_of(cursor, student_id):
    """`(department_code, program_code, year_level, section_name)` or None."""
    cursor.execute(
        """
        SELECT d.department_code, p.program_code, sec.year_level, sec.section_name
        FROM students s
        JOIN sections sec ON sec.id = s.section_id
        JOIN programs p ON p.id = sec.program_id
        JOIN departments d ON d.id = p.department_id
        WHERE s.student_id = %s
        """,
        (student_id,),
    )

    return cursor.fetchone()


def test_009_places_a_student_by_program_year_and_section(database_at_008):
    """
    The deployed shape exactly: the department text is `College of Computing`,
    which is not the name of any department, and the student is still placed -
    under the department their *program* belongs to.
    """
    cursor = database_at_008.cursor
    legacy_student(cursor, "INT-TEST-0001", "College of Computing", "BSCS", 4, "4B")
    database_at_008.migrate()

    assert placement_of(cursor, "INT-TEST-0001") == ("CCIT", "BSCS", 4, "4B")


def test_009_gives_students_in_one_section_one_section_row(database_at_008):
    """Three students, one section - not three sections of one student each."""
    cursor = database_at_008.cursor

    # Differing only in case and padding, which is how free text differs.
    legacy_student(cursor, "INT-TEST-0001", "College of Computing", "BSCS", 1, "B")
    legacy_student(cursor, "INT-TEST-0002", "College of Computing", "bscs", 1, "b")
    legacy_student(cursor, "INT-TEST-0003", "", " BSCS ", 1, " B ")
    database_at_008.migrate()

    cursor.execute("SELECT COUNT(*) FROM sections")
    assert cursor.fetchone()[0] == 1

    cursor.execute("SELECT COUNT(DISTINCT section_id), COUNT(section_id) FROM students")
    assert cursor.fetchone() == (1, 3)


@pytest.mark.parametrize(
    ("program", "year_level", "section"),
    [
        ("BSIS", 3, "A"),      # offered by the old dropdown, not a program
        ("BSCS", None, "A"),   # no year level
        ("BSCS", 3, ""),       # no section
        ("", None, ""),        # registered from the phone with nothing chosen
    ],
)
def test_009_leaves_a_student_it_cannot_place_unassigned_and_untouched(
    database_at_008, program, year_level, section
):
    """
    The data to be afraid of. Nothing is guessed, nothing fails, and - the
    part that makes "nothing is guessed" safe - what was typed is still there
    afterwards, because 009 drops no column.
    """
    cursor = database_at_008.cursor
    legacy_student(
        cursor, "INT-TEST-0001", "College of Computing", program, year_level, section
    )
    database_at_008.migrate()

    cursor.execute(
        """
        SELECT section_id IS NULL, college_department, program, year_level, section
        FROM students WHERE student_id = 'INT-TEST-0001'
        """
    )

    assert cursor.fetchone() == (
        1, "College of Computing", program, year_level, section
    )


def test_009_links_a_subject_to_the_account_its_typed_name_matches(database_at_008):
    """
    DM-4. On the deployed database every subject named an instructor in text
    and none was linked to the account of the same name, because 005's
    backfill ran when `instructors` was empty.
    """
    cursor = database_at_008.cursor
    account = make_instructor(cursor, fullname="Integration Instructor")
    subject = make_subject(cursor, instructor="Integration Instructor")
    stranger = make_subject(
        cursor, subject_code="INT-TEST-202", instructor="Somebody With No Account"
    )
    database_at_008.migrate()

    cursor.execute("SELECT id, instructor_id FROM subjects ORDER BY id")

    assert cursor.fetchall() == [(subject, account), (stranger, None)]


def test_009_does_not_link_a_name_two_accounts_share(database_at_008):
    """
    Two instructors with one name: linking by name would hand one of them the
    other's classes - and, now that reports are scoped by this key, the other's
    register. Left unlinked for a person to decide.
    """
    cursor = database_at_008.cursor
    make_instructor(cursor, "INT-TEST-INSTR-1", "Same Name")
    make_instructor(cursor, "INT-TEST-INSTR-2", "Same Name")
    make_subject(cursor, instructor="Same Name")
    database_at_008.migrate()

    cursor.execute("SELECT instructor_id FROM subjects")

    assert cursor.fetchone() == (None,)


def test_009_adopts_tables_that_were_created_by_hand(database_at_008):
    """
    DM-1. The deployed database already had `departments` and `programs`,
    created outside any migration, with its own rows and its own ids. 009 must
    run over that without failing, duplicating or renumbering.
    """
    cursor = database_at_008.cursor
    cursor.execute(
        """
        CREATE TABLE departments (
            id INT AUTO_INCREMENT PRIMARY KEY,
            department_code VARCHAR(50) NOT NULL,
            department_name VARCHAR(255) NOT NULL,
            UNIQUE KEY department_code (department_code),
            UNIQUE KEY department_name (department_name)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
        """
    )
    cursor.execute(
        """
        INSERT INTO departments (id, department_code, department_name)
        VALUES (41, 'CCIT', 'College of Computing and Information Technology')
        """
    )
    database_at_008.migrate()

    cursor.execute("SELECT id FROM departments WHERE department_code = 'CCIT'")
    assert cursor.fetchall() == [(41,)], "the existing row was replaced or duplicated"

    cursor.execute(
        "SELECT department_id FROM programs WHERE program_code = 'BSCS'"
    )
    assert cursor.fetchone() == (41,), "BSCS was not attached to the existing row"


def test_a_fresh_database_is_seeded_with_the_structure(scratch_database):
    """A new installation offers the same departments and programs."""
    from infra.db import db_cursor

    with db_cursor() as cursor:
        cursor.execute("SELECT COUNT(*) FROM departments")
        departments = cursor.fetchone()[0]

        cursor.execute(
            "SELECT COUNT(*) FROM programs p "
            "JOIN departments d ON d.id = p.department_id"
        )
        programs = cursor.fetchone()[0]

    assert (departments, programs) == (9, 9)


# ===========================================================================
# Placement, read back through the application
# ===========================================================================


def test_a_section_is_found_again_rather_than_created_twice(scratch_database):
    from infra.db import db_cursor
    from repositories import academics as academics_repo

    with db_cursor(commit=True) as cursor:
        program_id = academics_repo.program_id_for_code(cursor, "BSCS")

        first = academics_repo.section_for(cursor, program_id, 4, "B")
        again = academics_repo.section_for(cursor, program_id, 4, "b")
        other_year = academics_repo.section_for(cursor, program_id, 3, "B")

    assert first == again, "the same section, typed in another case, was duplicated"
    assert other_year != first, "a section is a year level as well as a name"


def test_a_section_cannot_be_created_under_a_program_that_does_not_exist(
    scratch_database,
):
    from infra.db import db_cursor
    from repositories import academics as academics_repo

    with db_cursor(commit=True) as cursor:
        assert academics_repo.section_for(cursor, 999999, 1, "A") is None

        cursor.execute("SELECT COUNT(*) FROM sections")
        assert cursor.fetchone()[0] == 0


def test_the_roster_reports_placement_under_the_old_key_names(scratch_database):
    """
    Templates and the mobile API read `college_department`, `program`,
    `year_level` and `section`. The storage changed; those keys did not.
    """
    from infra.db import db_cursor
    from repositories import students as students_repo

    with db_cursor(dictionary=True, commit=True) as cursor:
        make_student(cursor, "INT-TEST-0001", program="BSCS", year_level=4, section="B")
        make_student(cursor, "INT-TEST-0002", program=None)

        placed, unassigned = students_repo.all_students(cursor)

    assert (
        placed["college_department"],
        placed["program"],
        placed["year_level"],
        placed["section"],
    ) == ("College of Computing and Information Technology", "BSCS", 4, "B")

    assert unassigned["student_id"] == "INT-TEST-0002", "Unassigned sorts last"
    assert unassigned["section"] is None


def test_the_student_list_filters_by_each_level(scratch_database):
    from infra.db import db_cursor
    from repositories import students as students_repo

    with db_cursor(dictionary=True, commit=True) as cursor:
        make_student(cursor, "INT-TEST-0001", program="BSCS", year_level=4, section="B")
        make_student(cursor, "INT-TEST-0002", program="BSCS", year_level=1, section="B")
        make_student(cursor, "INT-TEST-0003", program="BSN", year_level=4, section="B")

        def ids(**filters):
            return [row["student_id"] for row in students_repo.roster(cursor, **filters)]

        bscs = students_repo.details(cursor, "INT-TEST-0001")

        assert ids(program_id=bscs["program_id"]) == ["INT-TEST-0002", "INT-TEST-0001"]
        assert ids(department_id=bscs["department_id"], year_level=4) == ["INT-TEST-0001"]
        assert ids(section_id=bscs["section_id"]) == ["INT-TEST-0001"]
        assert ids(year_level=4) == ["INT-TEST-0001", "INT-TEST-0003"]


def test_editing_a_student_moves_them_to_the_chosen_section(client, scratch_database):
    from infra.db import db_cursor
    from repositories import academics as academics_repo
    from repositories import students as students_repo

    with db_cursor(dictionary=True, commit=True) as cursor:
        make_student(cursor, "INT-TEST-0001", program="BSCS", year_level=1, section="B")
        nursing = academics_repo.program_id_for_code(cursor, "BSN")

    sign_in_as(client)

    response = client.post("/update_student/INT-TEST-0001", data={
        "name": "Integration Test Subject",
        "program_id": str(nursing),
        "year_level": "2",
        "section": "A",
    })

    assert response.status_code == 302

    with db_cursor(dictionary=True) as cursor:
        student = students_repo.details(cursor, "INT-TEST-0001")

    assert (student["program"], student["year_level"], student["section"]) == (
        "BSN", 2, "A"
    )


def test_an_instructor_is_listed_under_their_department(client, scratch_database):
    from infra.db import db_cursor
    from repositories import instructors as instructors_repo

    sign_in_as(client)

    with db_cursor(dictionary=True) as cursor:
        cursor.execute("SELECT id FROM departments WHERE department_code = 'CCIT'")
        department_id = cursor.fetchone()["id"]

    response = client.post("/add_instructor", data={
        "instructor_id": "INT-TEST-INSTR",
        "fullname": "Integration Instructor",
        "password": "a-password-nobody-uses",
        "department_id": str(department_id),
    })

    assert response.status_code == 302

    with db_cursor(dictionary=True) as cursor:
        (instructor,) = instructors_repo.roster(cursor, department_id=department_id)

    assert instructor["department_code"] == "CCIT"
    assert "password" not in instructor, "the hash reached a list screen (SE-17)"


# ===========================================================================
# DM-5 - an instructor sees their own classes and sessions
# ===========================================================================


@pytest.fixture
def two_instructors(scratch_database):
    """
    Two instructors, one subject each, one student marked in each, and one
    session each. Returns their `instructors.id`s and subject ids.
    """
    from infra.db import db_cursor

    with db_cursor(dictionary=True, commit=True) as cursor:
        mine = make_instructor(cursor, "INT-TEST-INSTR-1", "Instructor One")
        theirs = make_instructor(cursor, "INT-TEST-INSTR-2", "Instructor Two")

        my_subject = make_subject(
            cursor, subject_code="INT-TEST-101", instructor_id=mine
        )
        their_subject = make_subject(
            cursor, subject_code="INT-TEST-202", instructor_id=theirs
        )

        make_student(cursor, "INT-TEST-0001", name="Ana Reyes")
        make_student(cursor, "INT-TEST-0002", name="Ben Cruz")

        make_enrolment(cursor, "INT-TEST-0001", my_subject)
        make_enrolment(cursor, "INT-TEST-0002", their_subject)

        my_record = make_attendance(cursor, "INT-TEST-0001", my_subject)
        their_record = make_attendance(cursor, "INT-TEST-0002", their_subject)

        for subject_id in (my_subject, their_subject):
            cursor.execute(
                """
                INSERT INTO attendance_sessions
                    (subject_id, session_date, started_at, started_by)
                VALUES (%s, CURDATE(), NOW(), 'int-test')
                """,
                (subject_id,),
            )

    return {
        "mine": mine,
        "theirs": theirs,
        "my_subject": my_subject,
        "their_subject": their_subject,
        "my_record": my_record,
        "their_record": their_record,
    }


def page(client, url):
    response = client.get(url)

    assert response.status_code == 200

    return response.get_data(as_text=True)


def exported_students(client, url="/export_excel"):
    """The Student column of the exported sheet, without the header."""
    response = client.get(url)

    assert response.status_code == 200

    sheet = load_workbook(io.BytesIO(response.data)).active

    return [row[2] for row in sheet.iter_rows(min_row=2, values_only=True)]


def test_an_instructor_sees_only_their_own_register_and_sessions(
    client, two_instructors
):
    sign_in_as(client, "instructor", instructor_pk=two_instructors["mine"])

    html = page(client, "/reports")

    assert "Ana Reyes" in html and "INT-TEST-101" in html
    assert "Ben Cruz" not in html, "another instructor's student is on the page"
    assert "INT-TEST-202" not in html, (
        "another instructor's subject is on the page - in the register, the "
        "sessions table or the subject filter"
    )


def test_the_url_cannot_widen_an_instructors_scope(client, two_instructors):
    """
    The scope comes from the session. If `?instructor=` were read for an
    instructor it would be a default the URL could override; naming another
    instructor's subject must find nothing rather than find their register.
    """
    sign_in_as(client, "instructor", instructor_pk=two_instructors["mine"])

    theirs = two_instructors["theirs"]
    their_subject = two_instructors["their_subject"]

    assert "Ben Cruz" not in page(client, f"/reports?instructor={theirs}")
    assert "Ben Cruz" not in page(client, f"/reports?subject={their_subject}")

    assert exported_students(client, f"/export_excel?instructor={theirs}") == [
        "Ana Reyes"
    ]
    assert exported_students(client, f"/export_excel?subject={their_subject}") == []


def test_an_instructor_the_session_cannot_identify_sees_nothing(
    client, two_instructors
):
    """
    Fails closed. A session with no `instructor_pk` - one created before the
    key existed - must not be read as "unrestricted".
    """
    sign_in_as(client, "instructor")

    html = page(client, "/reports")

    assert "Ana Reyes" not in html and "Ben Cruz" not in html
    assert exported_students(client) == []


def test_the_administrator_sees_everything_or_one_instructor(
    client, two_instructors
):
    sign_in_as(client)

    everything = page(client, "/reports")

    assert "Ana Reyes" in everything and "Ben Cruz" in everything
    assert "Instructor One" in everything and "Instructor Two" in everything

    one = page(client, f"/reports?instructor={two_instructors['theirs']}")

    assert "Ben Cruz" in one and "Ana Reyes" not in one

    assert exported_students(
        client, f"/export_excel?instructor={two_instructors['theirs']}"
    ) == ["Ben Cruz"]
    assert sorted(exported_students(client)) == ["Ana Reyes", "Ben Cruz"]


def test_an_instructor_is_offered_only_their_own_subjects_to_run(
    client, two_instructors
):
    sign_in_as(client, "instructor", instructor_pk=two_instructors["mine"])

    html = page(client, "/attendance")

    assert "INT-TEST-101" in html
    assert "INT-TEST-202" not in html


def test_an_instructor_cannot_start_a_session_on_another_instructors_subject(
    client, two_instructors
):
    """
    The guard behind the filtered dropdown: a request that names the subject
    anyway. Refused before a session row exists and before the camera is
    touched.
    """
    from infra.db import db_cursor

    sign_in_as(client, "instructor", instructor_pk=two_instructors["mine"])

    response = client.post(
        "/start-attendance",
        data={"subject_id": str(two_instructors["their_subject"])},
    )

    body = response.get_json()

    assert body["success"] is False
    assert "not assigned to you" in body["message"]

    with db_cursor() as cursor:
        cursor.execute(
            "SELECT COUNT(*) FROM attendance_sessions WHERE subject_id = %s",
            (two_instructors["their_subject"],),
        )
        assert cursor.fetchone()[0] == 1, "a session row was opened for it"


def test_an_instructor_cannot_end_another_instructors_class(client, two_instructors):
    from infra.db import db_cursor

    sign_in_as(client, "instructor", instructor_pk=two_instructors["mine"])

    response = client.post(
        "/end-attendance",
        data={"subject_id": str(two_instructors["their_subject"])},
    )

    assert response.status_code == 403

    with db_cursor() as cursor:
        cursor.execute(
            "SELECT COUNT(*) FROM attendance_sessions "
            "WHERE subject_id = %s AND ended_at IS NOT NULL",
            (two_instructors["their_subject"],),
        )
        assert cursor.fetchone()[0] == 0, "their session was closed"


def test_an_instructor_cannot_correct_another_instructors_record(
    client, two_instructors
):
    from infra.db import db_cursor

    sign_in_as(client, "instructor", instructor_pk=two_instructors["mine"])

    their_record = two_instructors["their_record"]

    assert client.get(f"/attendance/{their_record}/correct").status_code == 404

    response = client.post(
        f"/attendance/{their_record}/correct",
        data={"status": "Absent", "reason": "Not mine to change"},
    )

    assert response.status_code == 404

    with db_cursor() as cursor:
        cursor.execute("SELECT status FROM attendance WHERE id = %s", (their_record,))
        assert cursor.fetchone()[0] == "Present", "the record was changed anyway"

        cursor.execute("SELECT COUNT(*) FROM attendance_audit")
        assert cursor.fetchone()[0] == 0

    # And their own still works - the refusal is not a blanket one.
    mine = client.post(
        f"/attendance/{two_instructors['my_record']}/correct",
        data={"status": "Late", "reason": "Arrived after the bell"},
    )

    assert mine.status_code == 302


def test_the_dashboard_counts_an_instructors_own_classes(client, two_instructors):
    from infra.db import db_cursor
    from repositories import attendance as attendance_repo

    with db_cursor(dictionary=True) as cursor:
        mine = attendance_repo.counters(
            cursor, instructor_id=two_instructors["mine"]
        )
        everyone = attendance_repo.counters(cursor)

    assert (
        mine["students"],
        mine["subjects"],
        mine["attendance_today"],
        mine["active_sessions"],
    ) == (1, 1, 1, 1)

    assert (everyone["students"], everyone["subjects"]) == (2, 2)


def test_a_session_that_recorded_nobody_is_still_listed(two_instructors):
    """
    The tallies are SUMs over a LEFT JOIN. A session with no attendance rows
    must appear with zeros, not vanish and not report NULL.
    """
    from infra.db import db_cursor
    from repositories import attendance as attendance_repo

    with db_cursor(dictionary=True) as cursor:
        (session_row,) = attendance_repo.sessions_filtered(
            cursor, instructor_id=two_instructors["mine"]
        )

    assert session_row["subject_code"] == "INT-TEST-101"
    assert (
        int(session_row["present"]),
        int(session_row["late"]),
        int(session_row["absent"]),
    ) == (0, 0, 0)


def test_a_subject_is_assigned_to_an_instructor_account(client, two_instructors):
    """DM-4. The form writes the key, so the offering belongs to somebody."""
    from infra.db import db_cursor

    sign_in_as(client)

    response = client.post("/add_subject", data={
        "subject_code": "INT-TEST-303",
        "subject_name": "Integration Testing",
        "instructor_id": str(two_instructors["mine"]),
        "day": "Monday",
        "course": "BSCS",
        "section": "B",
        "time_in": "08:00",
        "time_out": "10:00",
    })

    assert response.status_code == 302

    with db_cursor() as cursor:
        cursor.execute(
            "SELECT instructor_id FROM subjects WHERE subject_code = 'INT-TEST-303'"
        )
        assert cursor.fetchone()[0] == two_instructors["mine"]
